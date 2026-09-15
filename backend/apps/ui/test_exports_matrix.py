"""Parametrized export matrix (Phase 5): every document export family × format.

Each row asserts the endpoint returns 200, the right Content-Type, real file
magic bytes (so the byte stream is genuinely PDF/XLSX/CSV, not an error page),
and for XLSX that the workbook actually parses with a title cell present.

Covers the voucher exports added by the unification project (RFP / CV / PCF)
plus the JE single-entry downloads and the Engine-A format fallback.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.posting.models import PostingStatus

FORMATS = ["pdf", "xlsx", "csv"]
EXPECTED = {
    "pdf": ("application/pdf", b"%PDF-"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK\x03\x04"),
    "csv": ("text/csv", None),
}
PDF_MIN_BYTES = 800  # a rendered table PDF is never a stub/no-data page


def _assert_pdf_size(body):
    assert len(body) > PDF_MIN_BYTES, f"PDF is only {len(body)} bytes"


def _header_row(ws, first_label):
    """Locate the column-header row by its first cell (engine writes it once)."""
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == first_label and any(v not in (None, "") for v in row):
            return list(row)
    raise AssertionError(f"header starting {first_label!r} not found in workbook")


def _assert_money_numeric(ws, first_label, money_labels, marker=None):
    """Money columns (engine `money=True`) must land as numeric cells, never text.

    `marker` restricts the check to the data row that holds a known value
    (e.g. a seeded asset number) so the assertion runs against a real row.
    """
    header = _header_row(ws, first_label)
    missing = [label for label in money_labels if label not in header]
    assert not missing, f"money headers missing: {missing}"
    idx = [header.index(label) for label in money_labels]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    start = rows.index(header) + 1
    body = rows[start:]
    if marker is not None:
        body = [r for r in body if any(marker in str(c or "") for c in r)]
        assert body, f"no row contains marker {marker!r}"
        # regression guard: a data row must span columns horizontally, never
        # one-cell-per-row (the old Engine-A diagonal-layout bug).
        assert any(
            sum(1 for c in r if c not in (None, "")) >= 2 for r in body
        ), f"marker row {marker!r} collapsed to a single column"
    if not money_labels:
        return
    numerics = 0
    for r in body:
        for i in idx:
            v = r[i]
            if v in (None, ""):
                continue
            assert isinstance(v, (int, float, Decimal)), (
                f"{header[i]} cell {v!r} is text, not a number"
            )
            numerics += 1
    assert numerics > 0, "money columns exist but no numeric cell was found"


@pytest.fixture
def voucher_docs(db, company, segment, accounts, user):
    """One minimal RFP (+Dr/Cr lines), its CV, and a PCF replenishment."""
    from apps.ap.models import CheckVoucher, RFPDocument, RFPLine, Supplier
    from apps.cash.models import PCFReplenishment, PettyCashFund

    supplier = Supplier.objects.create(
        code="MX-01", name="Matrix Fuel Depot", default_segment=segment
    )
    rfp = RFPDocument.objects.create(
        ap_number="A9001",
        rfp_date=date(2026, 1, 7),
        payee=supplier,
        segment=segment,
        particulars="Matrix bulk fuel",
        purpose="purchase",
        amount=Decimal("100000.00"),
        status="prepared",
    )
    RFPLine.objects.create(
        rfp=rfp, line_no=1, side="dr", segment=segment,
        account=accounts["61100"], amount=Decimal("100000.00"),
        description="Matrix bulk fuel", cost_center="OS",
    )
    RFPLine.objects.create(
        rfp=rfp, line_no=2, side="cr", segment=segment,
        account=accounts["20000"], amount=Decimal("100000.00"),
        description="AP - Matrix Fuel Depot", cost_center="",
    )
    cv = CheckVoucher.objects.create(
        cv_number="CV-MX-0001",
        cv_date=date(2026, 1, 20),
        rfp=rfp,
        payee=supplier,
        bank_account=accounts["10110"],
        gross_amount=Decimal("100000.00"),
        withheld_tax=Decimal("0.00"),
        net_amount=Decimal("100000.00"),
        check_no="CHK-MX-1",
        status="created",
    )
    fund = PettyCashFund.objects.create(
        fund_code="PCF-MX", name="Matrix Fund", custodian=user,
        imprest_amount=Decimal("20000.00"), gl_account=accounts["10010"],
        company=company, is_active=True,
    )
    replen = PCFReplenishment.objects.create(
        fund=fund,
        request_date=date(2026, 1, 10),
        amount=Decimal("5000.00"),
        payee_name="Matrix Clerk",
        reference="PCF-MX-1",
        voucher_no="PCV-MX-0001",
        requested_by=user,
        customer_name="Matrix Customer",
        cost_center="OS",
        status="requested",
        expenses=[
            {
                "account_code": "61100", "amount": "5000.00", "side": "dr",
                "description": "Office supplies", "segment": "OS",
                "cost_center": "OS", "business_name": "Matrix Store", "tin": "000",
            }
        ],
    )
    return {"rfp": rfp, "cv": cv, "pcf": replen}


def _urls(docs):
    return {
        "rfp": f"/ap/rfps/{docs['rfp'].id}/export/",
        "cv": f"/ap/cv/{docs['cv'].id}/export/",
        "pcf": f"/cash/pcf/replenishments/{docs['pcf'].id}/export/",
    }


@pytest.mark.django_db
class TestVoucherExportMatrix:

    @pytest.mark.parametrize("family", ["rfp", "cv", "pcf"])
    @pytest.mark.parametrize("fmt", FORMATS)
    def test_export_responds_with_real_bytes(
        self, client, voucher_docs, user, family, fmt
    ):
        client.force_login(user)
        url = _urls(voucher_docs)[family] + f"{fmt}/"
        resp = client.get(url)
        assert resp.status_code == 200, f"{url} -> {resp.status_code}"

        content_type, magic = EXPECTED[fmt]
        assert resp["Content-Type"] == content_type
        body = resp.content
        if magic is not None:
            assert body[: len(magic)] == magic, f"{family}/{fmt} is not a real {fmt}"
        if fmt == "pdf":
            _assert_pdf_size(body)
        # attachments always carry a filename
        assert "attachment; filename=" in resp["Content-Disposition"]

    @pytest.mark.parametrize("family", ["rfp", "cv", "pcf"])
    def test_xlsx_parses_and_has_title(self, client, voucher_docs, user, family):
        client.force_login(user)
        url = _urls(voucher_docs)[family] + "xlsx/"
        resp = client.get(url)
        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        assert wb.sheetnames[0], "workbook must have a sheet"
        ws = wb.active
        assert ws["A1"].value, "title cell must be populated"
        # money columns must be numeric cells (RFP/CV "Amount", PCF Debit/Credit)
        first = "#" if family != "pcf" else "Business Name"
        money_labels = ["Debit", "Credit"] if family == "pcf" else ["Amount"]
        _assert_money_numeric(ws, first, money_labels)

    @pytest.mark.parametrize("family", ["rfp", "cv"])
    def test_pdf_is_voucher_form(self, client, voucher_docs, user, family):
        """The voucher PDFs must be the A5 form reproduction, not the table engine."""
        import re

        client.force_login(user)
        url = _urls(voucher_docs)[family] + "pdf/"
        resp = client.get(url)
        assert resp.status_code == 200
        body = resp.content
        mediabox = re.findall(rb"/MediaBox \[[^\]]*\]", body)
        assert mediabox, "PDF must declare a MediaBox"
        # A5 portrait (half-bond) — matches the print default.
        assert mediabox[0] == b"/MediaBox [ 0 0 419.5276 595.2756 ]", mediabox

    def test_bogus_format_falls_back_to_xlsx(self, client, voucher_docs, user):
        client.force_login(user)
        url = _urls(voucher_docs)["rfp"] + "bogus/"
        resp = client.get(url)
        assert resp.status_code == 200
        assert resp["Content-Type"] == EXPECTED["xlsx"][0]

    def test_csv_has_header_row(self, client, voucher_docs, user):
        client.force_login(user)
        for family in ("rfp", "cv", "pcf"):
            resp = client.get(_urls(voucher_docs)[family] + "csv/")
            assert resp.status_code == 200
            body = resp.content.decode("utf-8")
            assert "Account Name" in body or "GL Account" in body, family
            assert "Amount" in body or "Debit" in body, family


@pytest.mark.django_db
class TestConsoAssetReceiptMatrix:

    @pytest.fixture
    def docs(self, company, segment, accounts, user, voucher_docs):
        """A CONSO batch with one RFP member + a fixed asset with a schedule row."""
        from apps.ap.models import CONSOBatch, RFPDocument
        from apps.assets.models import Asset, AssetCategory
        from apps.assets.services import DepreciationService

        rfp = voucher_docs["rfp"]
        batch = CONSOBatch.objects.create(
            batch_no="CONSO-MX-01", conso_date=date(2026, 1, 15), status="open",
        )
        rfp.conso = batch
        rfp.save(update_fields=["conso", "updated_at"])
        batch.total_amount = rfp.amount
        batch.save(update_fields=["total_amount", "updated_at"])

        cat = AssetCategory.objects.create(
            code="VEHICLE", name="Vehicles", useful_life_years=5,
            asset_account=accounts["10010"],
            depreciation_expense_account=accounts["61100"],
            accumulated_dep_account=accounts["20000"],
        )
        asset = Asset.objects.create(
            asset_no="FA-MX-0001", name="Matrix Forklift", category=cat,
            segment=segment, acquisition_date=date(2026, 1, 5),
            cost=Decimal("100000.00"), residual_value=Decimal("5000.00"),
            asset_account=accounts["10010"],
            depreciation_expense_account=accounts["61100"],
            accumulated_dep_account=accounts["20000"],
            funding_source="cash", status="active",
        )
        DepreciationService.post_month(
            asset, period_start=date(2026, 2, 1), user=user,
        )
        return {"conso": batch, "asset": asset}

    @pytest.mark.parametrize("fmt", FORMATS)
    def test_conso_and_asset_exports(self, client, docs, user, fmt):
        client.force_login(user)
        content_type, magic = EXPECTED[fmt]
        for kind, obj in (("conso", docs["conso"]), ("asset", docs["asset"])):
            if kind == "conso":
                url = f"/ap/conso/{obj.id}/export/{fmt}/"
            else:
                url = f"/assets/{obj.id}/export/{fmt}/"
            resp = client.get(url)
            assert resp.status_code == 200, f"{url} -> {resp.status_code}"
            assert resp["Content-Type"] == content_type
            if magic is not None:
                assert resp.content[: len(magic)] == magic
            assert "attachment; filename=" in resp["Content-Disposition"]

    @pytest.mark.parametrize("fmt", FORMATS)
    def test_ar_receipts_export(self, client, voucher_docs, user, fmt):
        from apps.ar.models import AcknowledgmentReceipt, Customer

        customer = Customer.objects.create(
            code="MX-AR", name="Matrix Buyer", group="fuel",
            segment=voucher_docs["rfp"].segment, pricing_tier="regular",
        )
        ack = AcknowledgmentReceipt.objects.create(
            receipt_no="AR-MX-0001", customer=customer,
            transaction_date=date(2026, 1, 8), amount=Decimal("25000.00"),
            payment_method="cash", cash_account=voucher_docs["rfp"].lines.first().account,
            segment=voucher_docs["rfp"].segment,
        )
        client.force_login(user)
        for url in (f"/ar/receipts/export/{fmt}/",):
            resp = client.get(url)
            assert resp.status_code == 200, f"{url} -> {resp.status_code}"
        content_type, magic = EXPECTED[fmt]
        assert resp["Content-Type"] == content_type


@pytest.mark.django_db
class TestJeExportMatrix:

    @pytest.mark.parametrize("fmt", ["pdf", "xlsx", "csv"])
    def test_single_entry_exports(self, client, company, segment, accounts, user, fmt):
        from apps.ui.tests import _draft_entry

        je = _draft_entry(
            entry_no=f"MX-JE-{fmt}",
            transaction_date=date(2026, 1, 7),
            lines=[("61100", "85000.00"), ("20000", "-85000.00")],
            user=user,
        )
        client.force_login(user)
        assert je.status == PostingStatus.DRAFT
        resp = client.get(f"/journal/{je.id}/{fmt}/")
        assert resp.status_code == 200, f"/journal/{je.id}/{fmt}/"
        content_type, magic = EXPECTED[fmt]
        assert resp["Content-Type"] == content_type
        if magic is not None:
            assert resp.content[: len(magic)] == magic
            if fmt == "pdf":
                _assert_pdf_size(resp.content)
        if fmt == "xlsx":
            wb = load_workbook(BytesIO(resp.content))
            _assert_money_numeric(wb.active, "COA", ["Debit", "Credit"])


REGISTER_EXPORT_URLS = [
    "/ar/aging/export/",
    "/ap/aging/export/",
    "/ar/customers/export/",
    "/ap/suppliers/export/",
    "/foundation/coa/export/",
    "/assets/export/",
    "/cash/transfers/export/",
    "/cash/recon/export/",
    "/cash/short/export/",
    "/cash/collections-summary/export/",
    "/cash/collectibles/export/",
    "/cash/cycles/export/",
    "/cash/banks/export/",
    "/ap/advances/export/",
    "/ap/conso/export/",
    "/reports/fleet/fuel/export/",
    "/reports/tax/vat/export/",
    "/reports/tax/wht/export/",
    "/reports/tax/provision/export/",
    "/reports/tax/calendar/export/",
    "/reports/month-end-close/export/",
]

# First cell of each register's column-header row (engine writes it verbatim).
REGISTER_HEADERS = {
    "/ar/aging/export/": "Invoice #",
    "/ap/aging/export/": "AP #",
    "/ar/customers/export/": "Code",
    "/ap/suppliers/export/": "Code",
    "/foundation/coa/export/": "Code",
    "/assets/export/": "Asset No",
    "/cash/transfers/export/": "Date",
    "/cash/recon/export/": "Cycle",
    "/cash/short/export/": "Cycle",
    "/cash/collections-summary/export/": "DATE",
    "/cash/collectibles/export/": "Department",
    "/cash/cycles/export/": "Cycle start",
    "/cash/banks/export/": "Code",
    "/ap/advances/export/": "Employee",
    "/ap/conso/export/": "Batch No",
    "/reports/fleet/fuel/export/": "Date",
    "/reports/tax/vat/export/": "Invoice",
    "/reports/tax/wht/export/": "CV No",
    "/reports/tax/provision/export/": "Segment",
    "/reports/tax/calendar/export/": "Form",
    "/reports/month-end-close/export/": "Item",
}


@pytest.mark.django_db
class TestRegisterExportMatrix:
    """Smoke matrix for the Phase-3 register/list exports (Engine A, ?format=).

    Empty-data safe: every register tolerates zero rows, so each endpoint must
    return the format's artifact (real bytes) even on a fresh database.
    """

    @pytest.mark.parametrize("url", REGISTER_EXPORT_URLS)
    @pytest.mark.parametrize("fmt", FORMATS)
    def test_register_exports(self, client, user, url, fmt):
        client.force_login(user)
        resp = client.get(f"{url}?format={fmt}")
        assert resp.status_code == 200, f"{url}?format={fmt} -> {resp.status_code}"
        content_type, magic = EXPECTED[fmt]
        assert resp["Content-Type"] == content_type
        body = resp.content
        if magic is not None:
            assert body[: len(magic)] == magic
            if fmt == "pdf":
                _assert_pdf_size(body)
        assert "attachment; filename=" in resp["Content-Disposition"]
        if fmt == "xlsx":
            wb = load_workbook(BytesIO(body))
            ws = wb.active
            assert ws["A1"].value, f"{url}: title cell must be populated"
            assert any(
                row and row[0] == REGISTER_HEADERS[url]
                for row in ws.iter_rows(values_only=True)
            ), f"{url}: header row missing"
        elif fmt == "csv":
            first_line = body.decode("utf-8").splitlines()[0]
            assert first_line.strip(), f"{url}: CSV header must not be empty"
            assert REGISTER_HEADERS[url] in first_line

    def test_register_exports_require_login(self, client, url=REGISTER_EXPORT_URLS[0]):
        resp = client.get(f"{url}?format=xlsx")
        assert resp.status_code in (302, 403)


@pytest.mark.django_db
class TestRegisterDataQuality:
    """Data-backed register exports: a seeded row must actually appear, its
    money columns must be numeric cells in XLSX, the CSV must carry both the
    header and the seeded value, and the PDF must render a real document."""

    @pytest.fixture
    def register_data(self, company, segment, accounts, user):
        from apps.ap.models import Supplier
        from apps.ar.models import Customer
        from apps.assets.models import Asset, AssetCategory
        from apps.cash.models import CashShortExcessWorksheet, WeeklyCashCycle

        customer = Customer.objects.create(
            code="MX-100", name="Matrix Trading Corp", group="fuel",
            segment=segment, pricing_tier="regular",
        )
        supplier = Supplier.objects.create(
            code="MX-100", name="Matrix Supply Inc", default_segment=segment,
        )
        cat = AssetCategory.objects.create(
            code="EQ", name="Equipment", useful_life_years=5,
            asset_account=accounts["10010"],
            depreciation_expense_account=accounts["61100"],
            accumulated_dep_account=accounts["20000"],
        )
        asset = Asset.objects.create(
            asset_no="FA-MX-0002", name="Matrix Generator", category=cat,
            segment=segment, acquisition_date=date(2026, 2, 3),
            cost=Decimal("75000.00"), residual_value=Decimal("0.00"),
            asset_account=accounts["10010"],
            depreciation_expense_account=accounts["61100"],
            accumulated_dep_account=accounts["20000"],
            funding_source="cash", status="active",
        )
        cycle = WeeklyCashCycle.objects.create(
            cycle_start=date(2026, 3, 3), cycle_end=date(2026, 3, 9),
            segment=segment, closing_balance=Decimal("45678.90"),
        )
        CashShortExcessWorksheet.objects.create(
            cycle=cycle, segment=segment,
            expected_cash=Decimal("500.00"), actual_cash=Decimal("480.00"),
            variance=Decimal("20.00"), cause="Cashier typo", status="open",
        )
        return {"customer": customer, "supplier": supplier, "asset": asset}

    # url: (first header cell, money column labels, seeded value to find)
    DATA_URLS = {
        "/assets/export/": ("Asset No", ["Cost"], "FA-MX-0002"),
        "/cash/cycles/export/": ("Cycle start", ["Closing balance"], "2026-03-03"),
        "/cash/short/export/": ("Cycle", ["Expected", "Actual", "Variance"], "Cashier typo"),
        "/foundation/coa/export/": ("Code", [], "61100"),
        "/ar/customers/export/": ("Code", [], "Matrix Trading Corp"),
        "/ap/suppliers/export/": ("Code", [], "Matrix Supply Inc"),
    }

    @pytest.mark.parametrize("url", sorted(DATA_URLS))
    @pytest.mark.parametrize("fmt", FORMATS)
    def test_seeded_row_survives_export(self, client, register_data, user, url, fmt):
        client.force_login(user)
        resp = client.get(f"{url}?format={fmt}")
        assert resp.status_code == 200, f"{url}?format={fmt} -> {resp.status_code}"
        first, money_labels, marker = self.DATA_URLS[url]
        if fmt == "xlsx":
            ws = load_workbook(BytesIO(resp.content)).active
            _assert_money_numeric(ws, first, money_labels, marker=marker)
        elif fmt == "csv":
            body = resp.content.decode("utf-8")
            assert first in body.splitlines()[0], f"{url}: header row missing in CSV"
            assert marker in body, f"{url}: seeded value {marker!r} missing in CSV"
        else:
            _assert_pdf_size(resp.content)