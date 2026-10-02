"""Special Sales Invoice (Fuel Delivery) — UI coverage.

Covers the register (list/filter/HTMX), the create/edit form round-trip, the
draft → submitted → posted approval flow from the browser, rejection,
screen gating + nav + grant-editor registration, the My Approvals inbox
entry, and print/export downloads.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from openpyxl import load_workbook

from apps.ar.models import Customer, SpecialSalesInvoice
from apps.ar.services import SpecialInvoiceService
from apps.foundation.models import Account, UserProfile

pytestmark = pytest.mark.django_db

User = get_user_model()


@pytest.fixture
def fuel_sales_account(db, accounts):
    return Account.objects.create(
        code="40000",
        name="Sales - Fuel",
        account_type="revenue",
        segment=Account.segment_for_code("40000"),
    )


@pytest.fixture
def customer(db):
    return Customer.objects.create(code="C-SSI", name="SSI Fuel Client")


def _grid(segment, amount="15000.00"):
    return {
        "line_segment": [str(segment.id), str(segment.id)],
        "line_account": ["12030", "40000"],
        "line_debit": [amount, ""],
        "line_credit": ["", amount],
        "line_description": ["AR - fuel delivery", "Fuel sales"],
        "line_cost_center": ["GEN-FUEL", ""],
    }


def _clients(client, role_users):
    staff_client = client
    staff_client.force_login(role_users["staff"])
    head_client = type(client)()
    head_client.force_login(role_users["head"])
    return staff_client, head_client


def _post_new(staff_client, customer, segment, **overrides):
    data = {
        "customer": str(customer.id),
        "delivery_receipt_no": "DR-UI-1",
        "transaction_date": "2026-02-10",
        "segment": str(segment.id),
        "notes": "Tanker 7",
        **_grid(segment),
    }
    data.update(overrides)
    return staff_client.post("/ar/special-invoices/new/", data)


# ---------------------------------------------------------------------- pages


def test_ssi_pages_flow(client, customer, segment, fuel_sales_account, accounts, role_users):
    staff_client, head_client = _clients(client, role_users)

    resp = staff_client.get("/ar/special-invoices/new/")
    assert resp.status_code == 200
    body = resp.content.decode()
    assert "Account Distribution" in body
    assert "Delivery Receipt No." in body

    resp = _post_new(staff_client, customer, segment)
    assert resp.status_code == 302
    invoice = SpecialSalesInvoice.objects.latest("id")
    assert invoice.status == "draft"
    assert invoice.total == Decimal("15000.00")
    assert invoice.delivery_receipt_no == "DR-UI-1"
    assert invoice.invoice_no.startswith("SSI-2026-")

    resp = staff_client.get("/ar/special-invoices/")
    assert resp.status_code == 200
    assert invoice.invoice_no in resp.content.decode()

    detail = staff_client.get(f"/ar/special-invoices/{invoice.id}/")
    assert detail.status_code == 200
    content = detail.content.decode()
    assert invoice.invoice_no in content
    assert "DR-UI-1" in content
    assert "12030" in content and "40000" in content

    assert staff_client.get(f"/ar/special-invoices/{invoice.id}/print/").status_code == 200

    staff_client.post(f"/ar/special-invoices/{invoice.id}/submit/")
    head_client.post(f"/ar/special-invoices/{invoice.id}/approve/")
    invoice.refresh_from_db()
    assert invoice.status == "posted"
    assert invoice.journal_entry is not None
    assert invoice.journal_entry.is_posted

    # The detail page shows the document's own audit trail (created →
    # submitted → approved), not another document's history.
    detail = staff_client.get(f"/ar/special-invoices/{invoice.id}/")
    trail = detail.content.decode()
    assert "Submitted" in trail
    assert "Approved" in trail


def test_ssi_create_rejects_unbalanced(
    client, customer, segment, fuel_sales_account, accounts, role_users
):
    staff_client, _ = _clients(client, role_users)
    grid = _grid(segment)
    grid["line_credit"] = ["", "100.00"]  # Dr 15000 vs Cr 100
    resp = staff_client.post(
        "/ar/special-invoices/new/",
        {
            "customer": str(customer.id),
            "transaction_date": "2026-02-10",
            "segment": str(segment.id),
            **grid,
        },
    )
    assert resp.status_code == 200  # re-rendered with an error, never 500
    assert SpecialSalesInvoice.objects.count() == 0


def test_ssi_reject_from_ui(client, customer, segment, fuel_sales_account, accounts, role_users):
    staff_client, head_client = _clients(client, role_users)
    _post_new(staff_client, customer, segment)
    invoice = SpecialSalesInvoice.objects.latest("id")
    staff_client.post(f"/ar/special-invoices/{invoice.id}/submit/")
    head_client.post(f"/ar/special-invoices/{invoice.id}/reject/", {"note": "Wrong DR"})
    invoice.refresh_from_db()
    assert invoice.status == "draft"
    assert invoice.rejection_note == "Wrong DR"
    detail = staff_client.get(f"/ar/special-invoices/{invoice.id}/")
    assert "Revise &amp; resubmit" in detail.content.decode()


def test_ssi_non_preparer_submit_blocked(
    client, customer, segment, fuel_sales_account, accounts, role_users
):
    staff_client, _ = _clients(client, role_users)
    _post_new(staff_client, customer, segment)
    invoice = SpecialSalesInvoice.objects.latest("id")
    other = User.objects.create_user("other", password="x")
    UserProfile.objects.create(user=other, approval_role="staff")
    other_client = type(client)()
    other_client.force_login(other)
    other_client.post(f"/ar/special-invoices/{invoice.id}/submit/")
    invoice.refresh_from_db()
    assert invoice.status == "draft"


def test_ssi_edit_roundtrip(client, customer, segment, fuel_sales_account, accounts, role_users):
    staff_client, _ = _clients(client, role_users)
    _post_new(staff_client, customer, segment)
    invoice = SpecialSalesInvoice.objects.latest("id")
    resp = staff_client.get(f"/ar/special-invoices/{invoice.id}/edit/")
    assert resp.status_code == 200
    grid = _grid(segment, amount="9000.00")
    resp = staff_client.post(
        f"/ar/special-invoices/{invoice.id}/edit/",
        {
            "delivery_receipt_no": "DR-UI-2",
            "transaction_date": "2026-02-11",
            "segment": str(segment.id),
            "notes": "corrected",
            **grid,
        },
    )
    assert resp.status_code == 302
    invoice.refresh_from_db()
    assert invoice.total == Decimal("9000.00")
    assert invoice.delivery_receipt_no == "DR-UI-2"


def test_ssi_filter_and_htmx(client, customer, segment, fuel_sales_account, accounts, role_users, user):
    staff_client, _ = _clients(client, role_users)
    _post_new(staff_client, customer, segment)
    invoice = SpecialSalesInvoice.objects.latest("id")

    resp = staff_client.get("/ar/special-invoices/", {"status": "submitted"})
    assert resp.status_code == 200
    # The flash message names the invoice; the table row (detail link) is absent.
    assert f"/ar/special-invoices/{invoice.id}/" not in resp.content.decode()

    resp = staff_client.get("/ar/special-invoices/", {"q": "DR-UI-1"})
    assert f"/ar/special-invoices/{invoice.id}/" in resp.content.decode()

    resp = staff_client.get("/ar/special-invoices/", HTTP_HX_REQUEST="true")
    assert resp.status_code == 200
    assert f"/ar/special-invoices/{invoice.id}/" in resp.content.decode()


# ---------------------------------------------------------------------- exports


def _posted_ssi(customer, segment, fuel_sales_account, accounts, role_users, user):
    ssi = SpecialInvoiceService.create_ssi(
        customer=customer,
        transaction_date=date(2026, 2, 10),
        segment=segment,
        delivery_receipt_no="DR-EXP",
        lines=[
            {"account": "12030", "segment": segment, "debit": "15000.00",
             "description": "AR", "cost_center": "GEN-FUEL"},
            {"account": "40000", "segment": segment, "credit": "15000.00",
             "description": "Fuel sales"},
        ],
        created_by=user,
    )
    SpecialInvoiceService.submit(ssi, user=user)
    SpecialInvoiceService.approve(ssi, user=role_users["head"])
    return SpecialSalesInvoice.objects.get(pk=ssi.pk)


@pytest.mark.parametrize("fmt,content_type,magic", [
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK\x03\x04"),
    ("csv", "text/csv", None),
    ("pdf", "application/pdf", b"%PDF-"),
])
def test_ssi_export_formats(
    client, customer, segment, fuel_sales_account, accounts, role_users, user,
    fmt, content_type, magic,
):
    staff_client, _ = _clients(client, role_users)
    ssi = _posted_ssi(customer, segment, fuel_sales_account, accounts, role_users, user)
    resp = staff_client.get(f"/ar/special-invoices/{ssi.id}/export/{fmt}/")
    assert resp.status_code == 200
    assert resp["Content-Type"].startswith(content_type)
    body = resp.content
    if magic:
        assert body.startswith(magic)
    if fmt == "pdf":
        assert len(body) > 800
    if fmt == "xlsx":
        ws = load_workbook(BytesIO(body)).active
        text = " ".join(str(c or "") for row in ws.iter_rows(values_only=True) for c in row)
        assert "12030" in text and "40000" in text
        assert "Delivery Receipt No." in text
    if fmt == "csv":
        text = body.decode()
        assert ssi.invoice_no in text and "12030" in text


# ---------------------------------------------------------------------- access


def _narrow_user(username="narrow"):
    u = User.objects.create_user(username, password="x")
    UserProfile.objects.create(user=u, approval_role="staff", screen_access=["dashboard", "my_approvals"])
    return u


def test_ssi_screen_gating(client, customer, segment, fuel_sales_account, accounts, role_users):
    staff_client, head_client = _clients(client, role_users)
    assert staff_client.get("/ar/special-invoices/").status_code == 200
    assert head_client.get("/ar/special-invoices/").status_code == 200

    narrow_client = type(client)()
    narrow_client.force_login(_narrow_user())
    assert narrow_client.get("/ar/special-invoices/").status_code == 403
    assert narrow_client.get("/ar/special-invoices/new/").status_code == 403
    # Inbox approve stays reachable from My Approvals (approval actions gate
    # to the inbox, not the register).
    from apps.ui import screens as S

    assert S.screen_for_url_name("ssi_approve") == "my_approvals"
    assert S.screen_for_url_name("ssi_reject") == "my_approvals"


def test_ssi_in_sidebar_and_grant_editor(
    client, customer, segment, fuel_sales_account, accounts, role_users
):
    staff_client, head_client = _clients(client, role_users)
    sidebar = staff_client.get("/ar/special-invoices/").content.decode()
    assert "Special Sales Invoices (Fuel)" in sidebar

    narrow_client = type(client)()
    narrow_client.force_login(_narrow_user("narrow2"))
    assert "Special Sales Invoices (Fuel)" not in narrow_client.get("/").content.decode()

    form = head_client.get(f"/settings/users/{role_users['staff'].id}/update/").content.decode()
    assert 'value="ssi_list"' in form


def test_ssi_in_my_approvals(
    client, customer, segment, fuel_sales_account, accounts, role_users, user
):
    staff_client, head_client = _clients(client, role_users)
    ssi = SpecialInvoiceService.create_ssi(
        customer=customer,
        transaction_date=date(2026, 2, 10),
        segment=segment,
        lines=[
            {"account": "12030", "segment": segment, "debit": "500.00"},
            {"account": "40000", "segment": segment, "credit": "500.00"},
        ],
        created_by=role_users["staff"],
    )
    SpecialInvoiceService.submit(ssi, user=role_users["staff"])
    page = head_client.get("/approvals/").content.decode()
    assert ssi.invoice_no in page
    assert "Special Sales Invoices (Fuel)" in page
    assert f"/ar/special-invoices/{ssi.id}/reject/" in page
